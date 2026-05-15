from pathlib import Path
from docx import Document
import zipfile
import shutil


def extract_paragraphs(document: Document):
    """DOCX 일반 문단 텍스트 추출"""
    texts = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            texts.append(text)

    return texts


def extract_tables(document: Document):
    """DOCX 안의 실제 Word 표 텍스트 추출"""
    texts = []

    for table_index, table in enumerate(document.tables, start=1):
        texts.append(f"\n===== TABLE {table_index} START =====")

        for row in table.rows:
            row_texts = []

            for cell in row.cells:
                cell_text = cell.text.strip().replace("\n", " ")
                row_texts.append(cell_text)

            texts.append(" | ".join(row_texts))

        texts.append(f"===== TABLE {table_index} END =====\n")

    return texts


def extract_images_from_docx(docx_path: Path, image_output_dir: Path):
    """
    DOCX 내부 이미지 추출
    DOCX는 zip 구조이고 이미지는 word/media/ 안에 있음
    """
    image_output_dir.mkdir(parents=True, exist_ok=True)

    extracted_images = []

    with zipfile.ZipFile(docx_path, "r") as docx_zip:
        for file_name in docx_zip.namelist():
            if file_name.startswith("word/media/"):
                image_name = Path(file_name).name
                image_path = image_output_dir / image_name

                with docx_zip.open(file_name) as source:
                    with open(image_path, "wb") as target:
                        shutil.copyfileobj(source, target)

                extracted_images.append(image_path)

    return extracted_images


def load_docx(docx_path: Path):
    """DOCX 파일 열기"""
    return Document(docx_path)