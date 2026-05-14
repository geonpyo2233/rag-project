from pathlib import Path
from docx import Document


BASE_DIR = Path(__file__).resolve().parent.parent

DATA_DIR = BASE_DIR / "data"

OUTPUT_DIR = BASE_DIR / "output"
TEXT_DIR = OUTPUT_DIR / "text"

TEXT_DIR.mkdir(parents=True, exist_ok=True)


def extract_text_from_docx(docx_path: Path):
    document = Document(docx_path)

    texts = []

    print("\n===== DOCX TEXT EXTRACTION =====")

    # =========================
    # 일반 문단 추출
    # =========================
    print("\n[문단 추출]")

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()

        if text:
            texts.append(text)
            print(text)

    # =========================
    # 표(table) 추출
    # =========================
    print("\n[표 추출]")

    for table_index, table in enumerate(document.tables, start=1):
        texts.append(f"\n===== TABLE {table_index} START =====")

        print(f"\nTABLE {table_index}")

        for row in table.rows:
            row_texts = []

            for cell in row.cells:
                cell_text = cell.text.strip()

                # 줄바꿈 제거
                cell_text = cell_text.replace("\n", " ")

                row_texts.append(cell_text)

            row_result = " | ".join(row_texts)

            texts.append(row_result)

            print(row_result)

        texts.append(f"===== TABLE {table_index} END =====\n")

    result_text = "\n".join(texts)

    return result_text


def save_extracted_text(docx_path: Path, text: str):
    output_path = TEXT_DIR / f"{docx_path.stem}_direct_extract.txt"

    output_path.write_text(
        text,
        encoding="utf-8"
    )

    print("\n===== TEXT FILE SAVED =====")
    print(output_path)

    return output_path


def main():
    files = list(DATA_DIR.glob("*.docx"))

    if not files:
        raise FileNotFoundError("DOCX 파일이 없습니다.")

    docx_path = files[0]

    print("\n선택된 파일:")
    print(docx_path.name)

    extracted_text = extract_text_from_docx(docx_path)

    save_extracted_text(
        docx_path,
        extracted_text
    )


if __name__ == "__main__":
    main()