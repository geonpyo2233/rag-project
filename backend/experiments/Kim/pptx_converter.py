#python-pptx 구동용 import
from pptx import Presentation

#paddle pOCR 구동용 import
from paddleocr import PaddleOCR
from PIL import Image
import numpy as np
import io

#확인 후 불필요 시 제거
from pathlib import Path

# 하기 # 줄은 pptx 사용 예제입니다.
# text_runs will be populated with a list of strings,
#  one for each text run in presentation
#text_runs = []

#for slide in prs.slides:
    #for shape in slide.shapes:
    #    if not shape.has_text_frame:
    #        continue
    #    for paragraph in shape.text_frame.paragraphs:
    #        for run in paragraph.runs:
    #            text_runs.append(run.text)
    #            print(run.text);

#실제 코드 시작

#ocr 가동
ocr = PaddleOCR(
    use_doc_orientation_classify=False,
    use_doc_unwarping=False,
    use_textline_orientation=False,
    lang='korean',
    engine="paddle")

def extract_text_from_shape(shape, slide_items):

    # shape에서 텍스트를 추출해서 slide_items 리스트에 append.
    if hasattr(shape, "text") and shape.has_text_frame:
        for paragraph in shape.text_frame.paragraphs:
            text = paragraph.text.strip()
            if text:  # 공백만 있는 텍스트는 제외
                slide_items.append((shape.left, shape.top, text))
    
    #하고 싶은 것: 1. 사진인지 파악 2. OCR 돌려서 텍스트 변환 3. 텍스트 저장
    if shape.shape_type == 13 : #MSO_SHAPE.TYPE.PICTURE
        try:
            image = shape.image
            img = Image.open(io.BytesIO(image.blob))
            img = img.convert("RGB")
            img_np = np.array(img)
        
            output_dir = Path("./output")
            output_dir.mkdir(parents=True, exist_ok=True)
            preds = ocr.predict(img_np)

            for pred in preds:
                rec_texts = pred['rec_texts']

            texcon = ' '.join(rec_texts)

            #rec_texts = [pred['rec_texts'] for pred in preds]            
            #print(rec_texts)

            #검증용 코드
            #for res in preds:
                #res.print()
                #res.save_to_img(save_path=output_dir) ## Save the processed image
                #res.save_to_json(save_path=output_dir) ## Save the current image's structured result in JSON format
                
            slide_items.append((shape.left, shape.top, texcon))
        
        except Exception as e:
                print(f"OCR 실패: {e}")

    # 테이블 처리
    if shape.has_table:
        table = shape.table
        rows_text = []
        for row in table.rows:
            row_text = []
            for cell in row.cells:
                cell_text = []
                for paragraph in cell.text_frame.paragraphs:
                    text = paragraph.text.strip()
                    if text:
                        cell_text.append(text)
                row_text.append(' '.join(cell_text))  # 셀 내 텍스트를 결합
            rows_text.append(', '.join(row_text))  # 행 내 텍스트를 결합
        table_text = '\n'.join(rows_text)  # 테이블 전체 텍스트를 결합
        slide_items.append((shape.left, shape.top, table_text))

    # 그룹화된 개체 처리
    if shape.shape_type == 6:  # MSO_SHAPE_TYPE.GROUP
        for grouped_shape in shape.shapes:
            extract_text_from_shape(grouped_shape, slide_items)

def extract_text_ordered_by_position(pptx_file):
    prs = Presentation(pptx_file)

    slide_texts = {}

    for i, slide in enumerate(prs.slides, start=1):
        slide_items = []
        for shape in slide.shapes:
            extract_text_from_shape(shape, slide_items)

        # 좌상단에서 우하단 방향으로 정렬
        slide_items_sorted = sorted(slide_items, key=lambda x: (x[1], x[0]))
        slide_text = "\n".join(item[2] for item in slide_items_sorted)  # 줄바꿈 추가
        slide_texts[i] = slide_text

    return slide_texts

#pptx 파일 위치 지정 - 향후 연동 필요
pptx_file = 'test7.pptx'

slide_texts = {}

#슬라이드의 모든 문자열 긁어오기 후 파일명.txt 문서로 출력
ordered_text_per_slide = extract_text_ordered_by_position(pptx_file)
txt_file_name = pptx_file.rsplit('.', 1)[0] + '.txt'
if txt_file_name:
    with open(txt_file_name, "w", encoding="utf-8") as f:
        for slide_number in sorted(ordered_text_per_slide):
            slide_text = ordered_text_per_slide[slide_number]
            f.write(f"Slide {slide_number}:\n{slide_text}\n\n")

#작업 완료 확인 용
#print(txt_file_name + " 작업 완료")