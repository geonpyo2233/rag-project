from __future__ import annotations

from pathlib import Path


def parse_docx_text(source_path: Path) -> str:
    """Extract paragraph and table text from DOCX using python-docx."""
    from docx import Document

    source_path = source_path.resolve()
    if source_path.suffix.lower() != ".docx":
        raise ValueError("DOCX 파서는 .docx 파일만 처리합니다.")
    if not source_path.exists():
        raise FileNotFoundError(f"DOCX 파일을 찾을 수 없습니다: {source_path}")

    document = Document(str(source_path))
    lines: list[str] = []

    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if text:
            lines.append(text)

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
            if cells:
                lines.append("\t".join(cells))

    return "\n".join(lines).strip()
